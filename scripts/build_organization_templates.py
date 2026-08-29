from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "public" / "templates"

BLUE = "245FBE"
DARK_BLUE = "173252"
LIGHT_BLUE = "EAF2FF"
ORANGE = "EF9A52"
LIGHT_ORANGE = "FFF4E8"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "D9E2EC"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    # Thai is classified as a complex script in OOXML. Setting only eastAsia
    # makes some DOCX renderers drop Thai glyphs, so every script slot is set.
    run.font.name = "Tahoma"
    r_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{slot}"), "Tahoma")
    if size is not None:
        run.font.size = Pt(size)
        size_half_points = str(int(round(size * 2)))
        r_pr = run._element.get_or_add_rPr()
        size_cs = r_pr.find(qn("w:szCs"))
        if size_cs is None:
            size_cs = OxmlElement("w:szCs")
            r_pr.append(size_cs)
        size_cs.set(qn("w:val"), size_half_points)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("หน้า ")
    set_run_font(run, 9, color="5D6B7A")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Tahoma"
    for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
        normal._element.rPr.rFonts.set(qn(f"w:{slot}"), "Tahoma")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    normal.paragraph_format.line_spacing = 1.10

    heading_specs = {
        "Title": (24, BLUE, 12, 4),
        "Subtitle": (11, "496A8B", 0, 12),
        "Heading 1": (16, BLUE, 14, 6),
        "Heading 2": (13, BLUE, 12, 4),
        "Heading 3": (12, DARK_BLUE, 10, 3),
    }
    for style_name, (size, color, before, after) in heading_specs.items():
        style = doc.styles[style_name]
        style.font.name = "Tahoma"
        for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
            style._element.rPr.rFonts.set(qn(f"w:{slot}"), "Tahoma")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Tahoma"
        for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
            style._element.rPr.rFonts.set(qn(f"w:{slot}"), "Tahoma")
        style.font.size = Pt(11)
        style.font.color.rgb = RGBColor.from_string(DARK_BLUE)
        style.paragraph_format.left_indent = Inches(0.28)
        style.paragraph_format.first_line_indent = Inches(-0.16)
        style.paragraph_format.space_after = Pt(4)


def new_document(title: str, document_code: str, purpose: str) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)
    configure_styles(doc)

    props = doc.core_properties
    props.title = title
    props.subject = purpose
    props.author = ""
    props.last_modified_by = ""
    props.comments = ""
    props.keywords = ""
    props.category = ""

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = hp.add_run("PEOPLE PULSE  |  เอกสารแม่แบบองค์กร")
    set_run_font(run, 8.5, True, BLUE)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.text = "ฉบับร่างสำหรับปรับใช้ภายในองค์กร — ห้ามใช้โดยไม่ตรวจทาน"
    set_run_font(fp.runs[0], 8.5, False, "5D6B7A")
    add_page_number(footer.add_paragraph())

    masthead = doc.add_table(rows=1, cols=2)
    masthead.alignment = WD_TABLE_ALIGNMENT.CENTER
    masthead.autofit = False
    masthead.columns[0].width = Inches(4.85)
    masthead.columns[1].width = Inches(1.65)
    masthead.cell(0, 0).width = Inches(4.85)
    masthead.cell(0, 1).width = Inches(1.65)
    for cell in masthead.rows[0].cells:
        set_cell_margins(cell, 180, 180, 180, 180)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(masthead.cell(0, 0), BLUE)
    set_cell_shading(masthead.cell(0, 1), DARK_BLUE)
    left = masthead.cell(0, 0).paragraphs[0]
    left.paragraph_format.space_after = Pt(2)
    r = left.add_run(title)
    set_run_font(r, 18, True, WHITE)
    sub = masthead.cell(0, 0).add_paragraph(purpose)
    sub.paragraph_format.space_after = Pt(0)
    set_run_font(sub.add_run(""), 9.5, False, WHITE)
    if sub.runs:
        set_run_font(sub.runs[0], 9.5, False, WHITE)
    right = masthead.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = right.add_run("DRAFT\n")
    set_run_font(rr, 14, True, ORANGE)
    rr2 = right.add_run(f"รหัส {document_code}\nเวอร์ชัน [[VERSION]]")
    set_run_font(rr2, 9, True, WHITE)

    doc.add_paragraph()
    add_notice(
        doc,
        "คำเตือนก่อนนำไปใช้",
        "เอกสารนี้เป็นแม่แบบทั่วไป ไม่ใช่คำปรึกษากฎหมาย กรุณาให้ HR และที่ปรึกษากฎหมายตรวจความถูกต้อง ปรับให้ตรงกับข้อเท็จจริง นโยบายบริษัท และกฎหมายที่ใช้บังคับก่อนลงนามหรือออกเอกสารจริง",
    )
    return doc


def add_notice(doc: Document, title: str, body: str, fill: str = LIGHT_ORANGE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    cell = table.cell(0, 0)
    set_cell_margins(cell, 140, 180, 140, 180)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, 11, True, DARK_BLUE)
    p2 = cell.add_paragraph(body)
    p2.paragraph_format.space_after = Pt(0)
    if p2.runs:
        set_run_font(p2.runs[0], 10, False, DARK_BLUE)


def add_info_table(doc: Document, rows: Sequence[tuple[str, str]], columns: int = 2) -> None:
    pair_count = columns
    table = doc.add_table(rows=0, cols=pair_count * 2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    label_width = 1.12 if columns == 2 else 1.55
    value_width = 2.13 if columns == 2 else 4.95
    for idx in range(0, len(rows), pair_count):
        row = table.add_row()
        chunk = rows[idx : idx + pair_count]
        for pair_idx in range(pair_count):
            label_cell = row.cells[pair_idx * 2]
            value_cell = row.cells[pair_idx * 2 + 1]
            if pair_idx < len(chunk):
                label, value = chunk[pair_idx]
            else:
                label, value = "", ""
            label_cell.text = label
            value_cell.text = value
            set_cell_shading(label_cell, LIGHT_BLUE if label else WHITE)
            for cell in (label_cell, value_cell):
                set_cell_margins(cell)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if label_cell.paragraphs[0].runs:
                set_run_font(label_cell.paragraphs[0].runs[0], 9.5, bool(label), DARK_BLUE)
            if value_cell.paragraphs[0].runs:
                set_run_font(value_cell.paragraphs[0].runs[0], 10, False, DARK_BLUE)
            label_cell.width = Inches(label_width)
            value_cell.width = Inches(value_width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_paragraph(doc: Document, text: str, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, 11, True, DARK_BLUE)
        r2 = p.add_run(text[len(bold_prefix) :])
        set_run_font(r2, 11, False, DARK_BLUE)
    else:
        r = p.add_run(text)
        set_run_font(r, 11, False, DARK_BLUE)


def add_bullets(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(item)
        set_run_font(r, 11, False, DARK_BLUE)


def add_clause(doc: Document, number: str, title: str, paragraphs: Sequence[str], bullets: Sequence[str] = ()) -> None:
    add_heading(doc, f"{number} {title}", 2)
    for text in paragraphs:
        add_paragraph(doc, text)
    if bullets:
        add_bullets(doc, bullets)


def add_signature_table(doc: Document, parties: Sequence[tuple[str, str]]) -> None:
    add_heading(doc, "ส่วนลงนาม", 1)
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index in range(0, len(parties), 2):
        row = table.add_row()
        for col in range(2):
            cell = row.cells[col]
            set_cell_margins(cell, 140, 160, 140, 160)
            party = parties[index + col] if index + col < len(parties) else ("", "")
            if not party[0]:
                cell.text = ""
                continue
            role, extra = party
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run("ลงชื่อ __________________________________\n")
            set_run_font(r, 10.5, False, DARK_BLUE)
            r2 = p.add_run("( [[ชื่อ-นามสกุล]] )\n")
            set_run_font(r2, 10.5, True, DARK_BLUE)
            r3 = p.add_run(f"{role}\n{extra}\nวันที่ ____ / ____ / ______")
            set_run_font(r3, 9.5, False, "496A8B")
            cell.width = Inches(3.25)


def add_change_log(doc: Document) -> None:
    add_heading(doc, "บันทึกการควบคุมเอกสาร", 1)
    table = doc.add_table(rows=2, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    headers = ["เวอร์ชัน", "วันที่แก้ไข", "รายละเอียด", "ผู้อนุมัติ"]
    values = ["[[VERSION]]", "[[REVISION_DATE]]", "[[REVISION_NOTE]]", "[[APPROVER]]"]
    for c, text in enumerate(headers):
        table.cell(0, c).text = text
        set_cell_shading(table.cell(0, c), BLUE)
        set_cell_margins(table.cell(0, c))
        set_run_font(table.cell(0, c).paragraphs[0].runs[0], 9, True, WHITE)
    set_repeat_table_header(table.rows[0])
    for c, text in enumerate(values):
        table.cell(1, c).text = text
        set_cell_margins(table.cell(1, c))
        set_run_font(table.cell(1, c).paragraphs[0].runs[0], 9, False, DARK_BLUE)


def save(doc: Document, filename: str) -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_DIR / filename)


def build_lease() -> None:
    doc = new_document(
        "สัญญาเช่าสำนักงาน",
        "ORG-LEASE-001",
        "แม่แบบกำหนดสิทธิ หน้าที่ ค่าใช้จ่าย และการส่งมอบพื้นที่เช่า",
    )
    add_heading(doc, "ข้อมูลสัญญา", 1)
    add_info_table(
        doc,
        [
            ("วันที่ทำสัญญา", "[[AGREEMENT_DATE]]"),
            ("สถานที่", "[[AGREEMENT_PLACE]]"),
            ("ผู้ให้เช่า", "[[LESSOR_NAME]]"),
            ("เลขประจำตัว", "[[LESSOR_ID]]"),
            ("ผู้เช่า", "[[LESSEE_NAME]]"),
            ("เลขทะเบียน", "[[LESSEE_REGISTRATION_NO]]"),
            ("ผู้ติดต่อ", "[[CONTACT_NAME]]"),
            ("โทรศัพท์", "[[CONTACT_PHONE]]"),
        ],
    )
    add_clause(doc, "1.", "ทรัพย์สินและวัตถุประสงค์การเช่า", [
        "ผู้ให้เช่าตกลงให้เช่า และผู้เช่าตกลงเช่า พื้นที่ตั้งอยู่ที่ [[PREMISES_ADDRESS]] ขนาดประมาณ [[AREA_SQM]] ตารางเมตร พร้อมทรัพย์สิน/อุปกรณ์ตามบัญชีแนบท้าย [[ANNEX_NO]].",
        "ผู้เช่าจะใช้พื้นที่เพื่อ [[PERMITTED_USE]] เท่านั้น การเปลี่ยนวัตถุประสงค์หรือให้บุคคลอื่นใช้พื้นที่ต้องได้รับความยินยอมเป็นลายลักษณ์อักษรตามเงื่อนไขที่คู่สัญญาตกลง.",
    ])
    add_clause(doc, "2.", "ระยะเวลา การส่งมอบ และต่ออายุ", [
        "ระยะเวลาเช่าเริ่มวันที่ [[START_DATE]] ถึงวันที่ [[END_DATE]] รวม [[TERM_LENGTH]]. วันส่งมอบพื้นที่คือ [[HANDOVER_DATE]].",
        "การต่ออายุให้แจ้งล่วงหน้าไม่น้อยกว่า [[RENEWAL_NOTICE_DAYS]] วัน และมีผลเมื่อคู่สัญญาลงนามข้อตกลงใหม่หรือเอกสารแก้ไขเพิ่มเติมเท่านั้น.",
    ])
    add_clause(doc, "3.", "ค่าเช่า เงินประกัน ภาษี และค่าใช้จ่าย", [
        "ค่าเช่าเดือนละ [[MONTHLY_RENT]] บาท ชำระภายในวันที่ [[PAYMENT_DUE_DAY]] ของเดือน ผ่าน [[PAYMENT_METHOD]]. เงินประกันจำนวน [[SECURITY_DEPOSIT]] บาท.",
    ], [
        "ค่าไฟฟ้า น้ำประปา อินเทอร์เน็ต และบริการส่วนกลาง: [[UTILITY_ALLOCATION]]",
        "ภาษี ค่าธรรมเนียม และการหักภาษี ณ ที่จ่าย: [[TAX_ALLOCATION]]",
        "เงื่อนไขคืนเงินประกันและรายการหัก: [[DEPOSIT_RETURN_TERMS]]",
    ])
    add_clause(doc, "4.", "การใช้พื้นที่ ซ่อมบำรุง และการดัดแปลง", [
        "ผู้เช่าต้องดูแลพื้นที่ตามสมควร แจ้งความเสียหายโดยไม่ชักช้า และไม่ดัดแปลงโครงสร้าง ระบบไฟ ระบบอัคคีภัย หรือป้ายภายนอกโดยไม่ได้รับอนุมัติ.",
        "ผู้รับผิดชอบการซ่อมบำรุงประจำและการซ่อมโครงสร้างกำหนดดังนี้: [[MAINTENANCE_RESPONSIBILITY]].",
    ])
    add_clause(doc, "5.", "ประกันภัย ความปลอดภัย และการเข้าตรวจ", [
        "คู่สัญญาจะปฏิบัติตามมาตรการความปลอดภัย อาคาร และแผนฉุกเฉินที่ตกลง ผู้ให้เช่าอาจเข้าตรวจโดยแจ้งล่วงหน้า [[INSPECTION_NOTICE_HOURS]] ชั่วโมง เว้นแต่กรณีฉุกเฉิน.",
        "ข้อกำหนดประกันภัยและความรับผิด: [[INSURANCE_TERMS]].",
    ])
    add_clause(doc, "6.", "ผิดสัญญา การแก้ไข และเลิกสัญญา", [
        "เมื่อฝ่ายหนึ่งผิดสัญญา อีกฝ่ายจะแจ้งรายละเอียดและให้ระยะเวลาแก้ไข [[CURE_PERIOD_DAYS]] วัน เว้นแต่กรณีที่คู่สัญญาหรือกฎหมายกำหนดให้เลิกสัญญาได้ทันที.",
        "เงื่อนไขการบอกเลิกก่อนกำหนด การคืนพื้นที่ และค่าเสียหาย: [[EARLY_TERMINATION_TERMS]].",
    ])
    add_clause(doc, "7.", "การส่งหนังสือและการระงับข้อพิพาท", [
        "การแจ้งให้ส่งตามที่อยู่/อีเมลของคู่สัญญาที่ระบุไว้ การเปลี่ยนข้อมูลติดต่อมีผลเมื่อแจ้งเป็นลายลักษณ์อักษร.",
        "คู่สัญญาจะเจรจาโดยสุจริตก่อนใช้กระบวนการระงับข้อพิพาทตาม [[DISPUTE_PROCESS]] และกฎหมายที่ใช้บังคับ [[GOVERNING_LAW]].",
    ])
    add_heading(doc, "เอกสารแนบท้ายที่ควรจัดเตรียม", 1)
    add_bullets(doc, [
        "แผนผังและภาพถ่ายพื้นที่ ณ วันส่งมอบ",
        "บัญชีทรัพย์สิน มิเตอร์ กุญแจ และบัตรผ่าน",
        "หนังสือรับรอง/หลักฐานอำนาจลงนามของคู่สัญญา",
        "ตารางค่าใช้จ่ายและเงื่อนไขบริการส่วนกลาง",
    ])
    add_signature_table(doc, [
        ("ผู้ให้เช่า", "พยาน: [[LESSOR_WITNESS]]"),
        ("ผู้เช่า", "พยาน: [[LESSEE_WITNESS]]"),
    ])
    save(doc, "office-lease-agreement-template.docx")


def build_employment() -> None:
    doc = new_document(
        "สัญญาจ้างงาน",
        "HR-CONTRACT-001",
        "แม่แบบบันทึกเงื่อนไขการจ้าง หน้าที่ ผลตอบแทน และการคุ้มครองข้อมูล",
    )
    add_heading(doc, "ข้อมูลคู่สัญญาและตำแหน่ง", 1)
    add_info_table(doc, [
        ("นายจ้าง", "[[EMPLOYER_NAME]]"),
        ("เลขทะเบียน", "[[EMPLOYER_REGISTRATION_NO]]"),
        ("ลูกจ้าง", "[[EMPLOYEE_NAME]]"),
        ("เลขประจำตัว", "[[EMPLOYEE_ID_NO]]"),
        ("ตำแหน่ง", "[[POSITION_TITLE]]"),
        ("แผนก", "[[DEPARTMENT]]"),
        ("วันเริ่มงาน", "[[START_DATE]]"),
        ("สถานที่ทำงาน", "[[WORK_LOCATION]]"),
    ])
    add_clause(doc, "1.", "ลักษณะการจ้างและหน้าที่", [
        "นายจ้างตกลงจ้างลูกจ้างในลักษณะ [[EMPLOYMENT_TYPE]] ให้ปฏิบัติหน้าที่ตามคำบรรยายงานแนบท้าย รวมถึงงานที่เกี่ยวเนื่องและสมเหตุสมผลภายใต้ขอบเขตตำแหน่ง.",
        "ผู้บังคับบัญชาโดยตรงคือ [[MANAGER_TITLE]] เป้าหมายเริ่มต้นและวิธีประเมินผลระบุใน [[PERFORMANCE_PLAN_REFERENCE]].",
    ])
    add_clause(doc, "2.", "ระยะทดลองงานและการติดตามผล", [
        "ระยะทดลองงาน (ถ้ามี) ตั้งแต่ [[PROBATION_START]] ถึง [[PROBATION_END]] การยืนยันผลต้องทำเป็นลายลักษณ์อักษรและเป็นไปตามกฎหมาย/นโยบายที่ใช้บังคับ.",
    ])
    add_clause(doc, "3.", "ค่าตอบแทนและสิทธิประโยชน์", [
        "เงินเดือนพื้นฐาน [[BASE_SALARY]] บาทต่อ [[PAY_PERIOD]] จ่ายในวันที่ [[PAY_DATE]] ผ่าน [[PAYMENT_CHANNEL]]. ค่าคอมมิชชั่น/เงินเพิ่มทักษะ/โบนัส (ถ้ามี) เป็นไปตามเอกสาร [[COMPENSATION_POLICY_REFERENCE]].",
    ], [
        "เวลาทำงานและเวลาพัก: [[WORKING_HOURS]]",
        "วันหยุดและการลา: [[LEAVE_AND_HOLIDAY_POLICY]]",
        "สวัสดิการ: [[BENEFITS]]",
        "การทำงานล่วงเวลาและค่าตอบแทน: [[OVERTIME_TERMS]]",
    ])
    add_clause(doc, "4.", "ทรัพย์สิน ระบบ และความปลอดภัย", [
        "ลูกจ้างใช้ทรัพย์สิน บัญชี ระบบ และข้อมูลของนายจ้างเพื่อการทำงานตามสิทธิ์ที่ได้รับ ต้องปฏิบัติตามนโยบายความปลอดภัยสารสนเทศ และคืนทรัพย์สินเมื่อร้องขอหรือสิ้นสุดการจ้าง.",
    ])
    add_clause(doc, "5.", "ความลับ ข้อมูลส่วนบุคคล และผลงาน", [
        "การเข้าถึง ใช้ เปิดเผย และเก็บข้อมูลต้องเป็นไปตามวัตถุประสงค์งาน สิทธิ์ที่ได้รับ และเอกสาร [[PRIVACY_SECURITY_POLICY]]. เงื่อนไขความลับและสิทธิในผลงานให้ระบุอย่างชัดเจนในเอกสารแนบท้ายที่ผ่านการตรวจทาน.",
    ])
    add_clause(doc, "6.", "พัฒนา สกิล และการประเมิน", [
        "แผนพัฒนาสกิล การประเมิน KPI และเส้นทางเติบโตใช้เพื่อพัฒนางานตามเกณฑ์ที่ประกาศล่วงหน้า ผลประเมินไม่เปลี่ยนสิทธิค่าจ้างหรือวินัยโดยอัตโนมัตินอกเหนือจากกระบวนการที่ถูกต้อง.",
    ])
    add_clause(doc, "7.", "วินัย ข้อร้องทุกข์ และการสิ้นสุดการจ้าง", [
        "ลูกจ้างรับทราบช่องทางร้องทุกข์ [[GRIEVANCE_CHANNEL]] และข้อบังคับการทำงานฉบับ [[WORK_RULES_VERSION]]. การเตือน การลงโทษ หรือการสิ้นสุดการจ้างต้องพิจารณาข้อเท็จจริง เปิดโอกาสให้ชี้แจง และดำเนินการตามกฎหมาย/นโยบายที่ใช้บังคับ.",
    ])
    add_heading(doc, "เอกสารประกอบที่ควรแนบ", 1)
    add_bullets(doc, [
        "คำบรรยายงานและสายการบังคับบัญชา",
        "โครงสร้างค่าตอบแทนและเงื่อนไขค่าคอมมิชชั่น/เงินเพิ่ม",
        "ข้อบังคับการทำงานและนโยบายข้อมูลส่วนบุคคล",
        "แบบรับมอบทรัพย์สินและบัญชีระบบ",
    ])
    add_signature_table(doc, [
        ("ผู้มีอำนาจลงนามแทนนายจ้าง", "ตำแหน่ง [[SIGNATORY_TITLE]]"),
        ("ลูกจ้าง", "รับสำเนาแล้ว"),
    ])
    add_change_log(doc)
    save(doc, "employment-agreement-template.docx")


def build_nda() -> None:
    doc = new_document(
        "ข้อตกลงการรักษาความลับ (NDA)",
        "LEGAL-NDA-001",
        "แม่แบบกำหนดขอบเขตข้อมูลลับ การใช้ข้อมูล และการคืนหรือทำลายข้อมูล",
    )
    add_heading(doc, "ข้อมูลคู่สัญญา", 1)
    add_info_table(doc, [
        ("ผู้เปิดเผย", "[[DISCLOSING_PARTY]]"),
        ("ผู้รับข้อมูล", "[[RECEIVING_PARTY]]"),
        ("วัตถุประสงค์", "[[PERMITTED_PURPOSE]]"),
        ("วันที่มีผล", "[[EFFECTIVE_DATE]]"),
        ("ระยะเวลา", "[[AGREEMENT_TERM]]"),
        ("โครงการ/งาน", "[[PROJECT_NAME]]"),
    ])
    add_clause(doc, "1.", "ข้อมูลที่ได้รับการคุ้มครอง", [
        "“ข้อมูลลับ” หมายถึงข้อมูลที่ระบุว่าเป็นความลับหรือควรเข้าใจโดยสมเหตุสมผลว่าเป็นความลับ รวมถึงข้อมูลธุรกิจ ลูกค้า ราคา สูตร กระบวนการ ระบบ รหัสผ่าน แผนงาน และข้อมูลส่วนบุคคลที่เปิดเผยเพื่อวัตถุประสงค์ที่กำหนด.",
    ])
    add_clause(doc, "2.", "ข้อยกเว้น", [
        "ข้อมูลลับไม่รวมข้อมูลที่ผู้รับพิสูจน์ได้ว่าเป็นสาธารณะโดยมิใช่ความผิดของผู้รับ มีอยู่โดยชอบก่อนรับ เปิดเผยโดยบุคคลที่มีสิทธิ หรือผู้รับพัฒนาขึ้นอย่างอิสระโดยไม่ใช้ข้อมูลลับ.",
    ])
    add_clause(doc, "3.", "หน้าที่ของผู้รับข้อมูล", [
        "ผู้รับจะใช้ข้อมูลเฉพาะวัตถุประสงค์ที่อนุญาต เปิดเผยแก่ผู้ที่จำเป็นต้องรู้และมีหน้าที่คุ้มครองเทียบเท่า ใช้มาตรการรักษาความปลอดภัยตามความเสี่ยง และแจ้งเหตุสงสัยรั่วไหลผ่าน [[INCIDENT_CHANNEL]] โดยไม่ชักช้า.",
    ])
    add_clause(doc, "4.", "การใช้ AI และบริการภายนอก", [
        "ห้ามป้อนข้อมูลลับหรือข้อมูลส่วนบุคคลลงใน AI/บริการภายนอกที่องค์กรไม่ได้อนุมัติ ต้องปกปิดข้อมูลเท่าที่ทำได้ ใช้สิทธิ์เข้าถึงขั้นต่ำ และให้มนุษย์ตรวจผลก่อนนำไปใช้.",
    ])
    add_clause(doc, "5.", "การเปิดเผยตามกฎหมาย", [
        "หากจำเป็นต้องเปิดเผยตามกฎหมายหรือคำสั่งที่ชอบ ผู้รับจะแจ้งผู้เปิดเผยล่วงหน้าเท่าที่ทำได้ และจำกัดการเปิดเผยเฉพาะส่วนที่จำเป็น.",
    ])
    add_clause(doc, "6.", "คืน ทำลาย และเก็บรักษา", [
        "เมื่อสิ้นสุดวัตถุประสงค์หรือได้รับคำขอ ผู้รับจะคืนหรือลบข้อมูลและสำเนาตาม [[RETURN_DELETION_PROCESS]] เว้นแต่ต้องเก็บตามกฎหมาย/ระบบสำรองที่เข้าถึงอย่างจำกัด.",
    ])
    add_clause(doc, "7.", "ระยะเวลาคุ้มครองและการเยียวยา", [
        "หน้าที่รักษาความลับมีผลเป็นเวลา [[CONFIDENTIALITY_PERIOD]] หลังการเปิดเผยหรือสิ้นสุดความสัมพันธ์ โดยให้ที่ปรึกษากฎหมายกำหนดระยะเวลาและมาตรการเยียวยาที่เหมาะสมกับประเภทข้อมูล.",
    ])
    add_signature_table(doc, [
        ("ผู้เปิดเผยข้อมูล", "ตำแหน่ง [[DISCLOSER_TITLE]]"),
        ("ผู้รับข้อมูล", "ตำแหน่ง [[RECIPIENT_TITLE]]"),
    ])
    add_change_log(doc)
    save(doc, "confidentiality-nda-template.docx")


def build_warning() -> None:
    doc = new_document(
        "หนังสือเตือนพนักงาน",
        "HR-WARNING-001",
        "แม่แบบบันทึกข้อเท็จจริง คำชี้แจง แผนแก้ไข และสิทธิอุทธรณ์อย่างเป็นธรรม",
    )
    add_notice(
        doc,
        "หลักสำคัญ",
        "หนังสือนี้ต้องออกจากข้อเท็จจริงที่ตรวจสอบแล้ว เปิดโอกาสให้พนักงานชี้แจง และใช้เกณฑ์เดียวกันอย่างไม่เลือกปฏิบัติ การลงชื่อ “รับทราบ” ไม่เท่ากับ “ยอมรับผิด” และการออกใบเตือนไม่ควรหัก Points อัตโนมัติ",
        fill=LIGHT_BLUE,
    )
    add_heading(doc, "ข้อมูลเอกสารและผู้เกี่ยวข้อง", 1)
    add_info_table(doc, [
        ("เลขที่เอกสาร", "[[WARNING_NO]]"),
        ("ระดับ", "[[WARNING_LEVEL]]"),
        ("ชื่อพนักงาน", "[[EMPLOYEE_NAME]]"),
        ("รหัสพนักงาน", "[[EMPLOYEE_ID]]"),
        ("ตำแหน่ง", "[[POSITION_TITLE]]"),
        ("แผนก", "[[DEPARTMENT]]"),
        ("วันที่เกิดเหตุ", "[[INCIDENT_DATE]]"),
        ("วันที่ออกเอกสาร", "[[ISSUED_DATE]]"),
        ("ผู้ออกเอกสาร", "[[ISSUED_BY]]"),
        ("ผู้ตรวจทาน HR", "[[HR_REVIEWER]]"),
    ])
    add_heading(doc, "1. ข้อเท็จจริงที่ตรวจสอบแล้ว", 2)
    add_paragraph(doc, "[[FACTUAL_DESCRIPTION: ระบุวัน เวลา สถานที่ เหตุการณ์ และข้อมูลที่ตรวจสอบได้ หลีกเลี่ยงคำตัดสินนิสัยหรือถ้อยคำดูหมิ่น]]")
    add_heading(doc, "2. กฎ/หน้าที่ที่เกี่ยวข้อง", 2)
    add_info_table(doc, [
        ("เอกสารอ้างอิง", "[[POLICY_OR_DUTY_REFERENCE]]"),
        ("เวอร์ชัน/วันที่มีผล", "[[REFERENCE_VERSION]]"),
    ], columns=1)
    add_heading(doc, "3. หลักฐานที่ใช้พิจารณา", 2)
    add_bullets(doc, [
        "[[EVIDENCE_1]]",
        "[[EVIDENCE_2]]",
        "[[EVIDENCE_3]]",
    ])
    add_heading(doc, "4. คำชี้แจงของพนักงาน", 2)
    add_paragraph(doc, "[[EMPLOYEE_EXPLANATION: บันทึกคำชี้แจง เหตุจำเป็น หรือข้อโต้แย้งตามที่พนักงานให้ข้อมูล]]")
    add_heading(doc, "5. ข้อสรุปและแผนแก้ไข", 2)
    add_info_table(doc, [
        ("ข้อสรุป", "[[FINDING]]"),
        ("สิ่งที่ต้องปรับปรุง", "[[CORRECTIVE_ACTION]]"),
        ("ผู้สนับสนุน", "[[SUPPORT_OWNER]]"),
        ("กำหนดเสร็จ", "[[TARGET_DATE]]"),
        ("วันติดตามผล", "[[REVIEW_DATE]]"),
        ("ผลหากเกิดซ้ำ", "[[CONSEQUENCE_IF_REPEATED]]"),
    ], columns=1)
    add_heading(doc, "6. สิทธิชี้แจงและอุทธรณ์", 2)
    add_paragraph(doc, "พนักงานสามารถส่งคำชี้แจง/อุทธรณ์ผ่าน [[APPEAL_CHANNEL]] ภายใน [[APPEAL_DAYS]] วัน โดยผู้ตรวจที่ไม่ใช่คู่กรณีจะตอบภายใน [[RESPONSE_SLA]] วันทำการ เอกสารและข้อมูลจะถูกจำกัดสิทธิ์ตามหน้าที่.")
    add_heading(doc, "การรับทราบ", 1)
    add_notice(doc, "ข้อความรับทราบ", "ข้าพเจ้าได้รับสำเนาและรับทราบว่าเอกสารนี้ถูกแจ้งแก่ข้าพเจ้าแล้ว การลงชื่อรับทราบไม่ถือเป็นการสละสิทธิชี้แจง อุทธรณ์ หรือยอมรับข้อกล่าวหา", fill=LIGHT_GRAY)
    add_signature_table(doc, [
        ("พนักงาน — รับทราบการรับเอกสาร", "หมายเหตุ [[EMPLOYEE_NOTE]]"),
        ("ผู้บังคับบัญชา/ผู้ออกเอกสาร", "ตำแหน่ง [[ISSUER_TITLE]]"),
        ("HR/ผู้ตรวจทาน", "สถานะ [[REVIEW_STATUS]]"),
        ("พยานการส่งมอบ (ถ้ามี)", "เหตุผลเมื่อพนักงานไม่ลงชื่อ [[DELIVERY_NOTE]]"),
    ])
    save(doc, "employee-warning-letter-template.docx")


def build_asset_handover() -> None:
    doc = new_document(
        "แบบรับมอบและคืนทรัพย์สิน",
        "OPS-ASSET-001",
        "แม่แบบบันทึกอุปกรณ์ บัญชีระบบ สภาพ และการคืนทรัพย์สินอย่างตรวจสอบได้",
    )
    add_heading(doc, "ข้อมูลผู้รับมอบ", 1)
    add_info_table(doc, [
        ("ชื่อ", "[[EMPLOYEE_NAME]]"),
        ("รหัสพนักงาน", "[[EMPLOYEE_ID]]"),
        ("ตำแหน่ง", "[[POSITION_TITLE]]"),
        ("แผนก", "[[DEPARTMENT]]"),
        ("ประเภทการทำรายการ", "[[HANDOVER_OR_RETURN]]"),
        ("วันที่", "[[TRANSACTION_DATE]]"),
        ("ผู้ส่งมอบ", "[[HANDOVER_OWNER]]"),
        ("สถานที่", "[[LOCATION]]"),
    ])
    add_heading(doc, "รายการทรัพย์สิน", 1)
    table = doc.add_table(rows=6, cols=7)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    headers = ["ลำดับ", "ทรัพย์สิน", "รหัส/Serial", "จำนวน", "สภาพ", "อุปกรณ์เสริม", "หมายเหตุ"]
    widths = [0.45, 1.25, 1.2, 0.55, 0.8, 1.1, 1.15]
    for col, (header, width) in enumerate(zip(headers, widths)):
        table.cell(0, col).text = header
        table.cell(0, col).width = Inches(width)
        set_cell_shading(table.cell(0, col), BLUE)
        set_cell_margins(table.cell(0, col), 80, 70, 80, 70)
        set_run_font(table.cell(0, col).paragraphs[0].runs[0], 8.5, True, WHITE)
    set_repeat_table_header(table.rows[0])
    for row in range(1, 6):
        values = [str(row), "[[ASSET_NAME]]", "[[ASSET_NO]]", "[[QTY]]", "[[CONDITION]]", "[[ACCESSORY]]", "[[NOTE]]"]
        for col, value in enumerate(values):
            table.cell(row, col).text = value
            table.cell(row, col).width = Inches(widths[col])
            set_cell_margins(table.cell(row, col), 90, 70, 90, 70)
            set_run_font(table.cell(row, col).paragraphs[0].runs[0], 8.5, False, DARK_BLUE)

    add_heading(doc, "บัญชีและสิทธิ์ระบบ", 1)
    add_bullets(doc, [
        "บัญชี/ระบบที่เปิดหรือปิดสิทธิ์: [[SYSTEM_ACCESS_LIST]]",
        "อุปกรณ์ยืนยันตัวตน/บัตรผ่าน: [[ACCESS_TOKEN_OR_CARD]]",
        "ผู้อนุมัติสิทธิ์และวันที่: [[ACCESS_APPROVER_AND_DATE]]",
        "ห้ามบันทึกรหัสผ่านจริงลงในแบบฟอร์มนี้",
    ])
    add_heading(doc, "เงื่อนไขการดูแลและคืน", 1)
    add_bullets(doc, [
        "ใช้เพื่อการทำงานตามสิทธิ์ที่ได้รับ ดูแลตามสมควร และแจ้งสูญหาย/เสียหายทันทีผ่าน [[INCIDENT_CHANNEL]]",
        "ไม่ติดตั้งซอฟต์แวร์หรือให้บุคคลอื่นใช้โดยไม่ได้รับอนุญาต",
        "คืนทรัพย์สิน อุปกรณ์เสริม ข้อมูล และสิทธิ์เข้าถึงเมื่อมีคำขอ ย้ายงาน หรือสิ้นสุดการจ้าง",
        "การเรียกเก็บความเสียหายต้องตรวจข้อเท็จจริงและดำเนินการตามกฎหมาย/นโยบาย ไม่หักเงินหรือ Points อัตโนมัติ",
    ])
    add_heading(doc, "ผลการตรวจรับ/คืน", 1)
    add_info_table(doc, [
        ("ผลตรวจ", "[[INSPECTION_RESULT]]"),
        ("รายการค้าง", "[[OUTSTANDING_ITEMS]]"),
        ("กำหนดแก้ไข", "[[RESOLUTION_DUE_DATE]]"),
        ("หลักฐานแนบ", "[[ATTACHMENTS]]"),
    ], columns=1)
    add_signature_table(doc, [
        ("ผู้รับมอบ/ผู้คืน", "ยืนยันรายการตามที่บันทึก"),
        ("ผู้ส่งมอบ/ผู้ตรวจรับ", "แผนก [[OWNER_DEPARTMENT]]"),
    ])
    add_change_log(doc)
    save(doc, "asset-handover-return-template.docx")


def main() -> None:
    build_lease()
    build_employment()
    build_nda()
    build_warning()
    build_asset_handover()
    for path in sorted(OUTPUT_DIR.glob("*.docx")):
        print(path)


if __name__ == "__main__":
    main()
